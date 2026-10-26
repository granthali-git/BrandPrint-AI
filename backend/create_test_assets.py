"""Script to generate test assets (.pptx, .pdf, .docx) for document parser verification."""
import os
import io

def generate_assets():
    assets_dir = os.path.join(os.path.dirname(__file__), "test_assets")
    os.makedirs(assets_dir, exist_ok=True)

    # 1. Generate 3-slide PPTX
    import pptx
    from pptx.util import Inches, Pt
    prs = pptx.Presentation()
    blank_layout = prs.slide_layouts[6]

    # Slide 1
    s1 = prs.slides.add_slide(blank_layout)
    tb1 = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    p1 = tb1.text_frame.paragraphs[0]
    p1.text = "AstraForge Cloud Overview"
    p1.font.size = Pt(28)

    tb1_body = s1.shapes.add_textbox(Inches(1), Inches(2.2), Inches(8), Inches(3))
    tf1 = tb1_body.text_frame
    b1 = tf1.paragraphs[0]
    b1.text = "Enterprise cloud scalability verified across global regions"
    b1.level = 0
    b2 = tf1.add_paragraph()
    b2.text = "High-availability containerized microservices architecture"
    b2.level = 1

    s1.notes_slide.notes_text_frame.text = "Speaker Note: Highlight our multi-region architecture."

    # Slide 2
    s2 = prs.slides.add_slide(blank_layout)
    tb2 = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = "Security & Compliance Audits"
    p2.font.size = Pt(28)

    table_shape = s2.shapes.add_table(2, 2, Inches(1), Inches(2.2), Inches(6), Inches(1.5))
    t = table_shape.table
    t.cell(0, 0).text = "Certification"
    t.cell(0, 1).text = "Audit Status"
    t.cell(1, 0).text = "SOC-2 Type II"
    t.cell(1, 1).text = "Fully Verified & Compliant"

    # Slide 3
    s3 = prs.slides.add_slide(blank_layout)
    tb3 = s3.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    p3 = tb3.text_frame.paragraphs[0]
    p3.text = "Enterprise SLA Commitments"
    p3.font.size = Pt(28)

    tb3_body = s3.shapes.add_textbox(Inches(1), Inches(2.2), Inches(8), Inches(3))
    tf3 = tb3_body.text_frame
    b3 = tf3.paragraphs[0]
    b3.text = "Our infrastructure provides 99.95% verified uptime backed by public status dashboards."
    b4 = tf3.add_paragraph()
    b4.text = "Backed by comprehensive SOC-2 Type II third-party audit reports."

    pptx_path = os.path.join(assets_dir, "enterprise_presentation.pptx")
    prs.save(pptx_path)
    print(f"Generated PPTX: {pptx_path}")

    # 2. Generate Multi-Page PDF
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    pdf_path = os.path.join(assets_dir, "compliance_whitepaper.pdf")
    c = canvas.Canvas(pdf_path, pagesize=letter)
    
    # Page 1
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, 720, "AstraForge Cloud Architecture Whitepaper")
    c.setFont("Helvetica", 12)
    c.drawString(72, 680, "Page 1: System Scalability and Multi-Tenant Isolation")
    c.drawString(72, 650, "AstraForge Enterprise Cloud provides resilient infrastructure for enterprise teams.")
    c.drawString(72, 620, "Our architecture guarantees strict boundary isolation across all tenant services.")
    c.showPage()

    # Page 2
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, 720, "Security & Regulatory Compliance Standards")
    c.setFont("Helvetica", 12)
    c.drawString(72, 680, "Page 2: Verified Uptime and Audit Certifications")
    c.drawString(72, 650, "Our infrastructure provides 99.95% verified uptime backed by real-time status dashboards.")
    c.drawString(72, 620, "All customer data is encrypted in transit and at rest adhering to SOC-2 Type II standards.")
    c.showPage()
    c.save()
    print(f"Generated PDF: {pdf_path}")

    # 3. Generate Word DOCX
    import docx
    doc = docx.Document()
    doc.add_heading("AstraForge Brand Policy & Technical Specification", level=1)
    doc.add_paragraph("This document outlines official corporate brand and service level specifications.")
    doc.add_heading("High-Availability Commitments", level=2)
    doc.add_paragraph("Our infrastructure provides 99.95% verified uptime backed by real-time public status dashboards.")
    
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Standard"
    table.rows[0].cells[1].text = "Level"
    table.rows[1].cells[0].text = "Availability SLA"
    table.rows[1].cells[1].text = "99.95% High Availability"

    docx_path = os.path.join(assets_dir, "product_guidelines.docx")
    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")

if __name__ == "__main__":
    generate_assets()
