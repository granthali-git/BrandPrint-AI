"""Universal document parser for enterprise compliance pipelines.
Extracts clean text, section headings, bullet points, speaker notes, and slide hierarchies
from PDF, DOCX, PPTX, and plaintext files.
"""
import io
import os
import re
import logging
from typing import Optional, List, Dict, Any
from fastapi import HTTPException, status

logger = logging.getLogger(__name__)

# Supported image and document extensions
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".doc",
    ".pptx",
    ".ppt",
    ".txt",
    ".md",
    ".csv",
    ".json",
    *IMAGE_EXTENSIONS,
}

MAX_WORD_LIMIT = 25000
MAX_IMAGE_FILE_SIZE = 10 * 1024 * 1024  # 10 MB per guideline image


def validate_and_truncate_text(text: str, max_words: int = MAX_WORD_LIMIT) -> str:
    """Validate text length and safely truncate if exceeding context limit."""
    if not text:
        return ""
    words = text.split()
    if len(words) > max_words:
        truncated_words = words[:max_words]
        logger.warning(f"Extracted document text exceeded {max_words} words ({len(words)} words). Truncating.")
        return " ".join(truncated_words) + f"\n\n[Note: Document exceeded maximum limit of {max_words:,} words. Content was safely truncated for compliance analysis.]"
    return text


def extract_from_pdf(file_bytes: bytes, filename: str) -> str:
    """Extract text from PDF pages with clear delimiters and graceful scanned-page handling."""
    if not file_bytes:
        return ""

    file_stream = io.BytesIO(file_bytes)
    pages_text: List[str] = []

    # 1. Attempt extraction using pypdf
    try:
        import pypdf
        reader = pypdf.PdfReader(file_stream)
        if reader.is_encrypted:
            try:
                # Attempt default empty password decrypt
                reader.decrypt("")
            except Exception:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"The uploaded PDF '{filename}' is password-protected. Please remove password encryption before uploading."
                )

        for page_idx, page in enumerate(reader.pages, start=1):
            try:
                page_raw = page.extract_text() or ""
                page_clean = page_raw.strip()
                if page_clean:
                    pages_text.append(f"--- [Page {page_idx}] ---\n\n{page_clean}")
                else:
                    pages_text.append(f"--- [Page {page_idx}] ---\n\n[Empty or Scanned Page]")
            except Exception as page_err:
                logger.warning(f"Error reading page {page_idx} of '{filename}' via pypdf: {page_err}")
                pages_text.append(f"--- [Page {page_idx}] ---\n\n[Page Content Unreadable]")

        if pages_text:
            return "\n\n".join(pages_text).strip()

    except HTTPException:
        raise
    except Exception as e:
        logger.warning(f"pypdf extraction failed on '{filename}': {e}. Attempting fallback to pdfplumber.")

    # 2. Fallback extraction using pdfplumber
    try:
        import pdfplumber
        file_stream.seek(0)
        pages_text = []
        with pdfplumber.open(file_stream) as pdf:
            for page_idx, page in enumerate(pdf.pages, start=1):
                try:
                    page_raw = page.extract_text() or ""
                    page_clean = page_raw.strip()
                    if page_clean:
                        pages_text.append(f"--- [Page {page_idx}] ---\n\n{page_clean}")
                    else:
                        pages_text.append(f"--- [Page {page_idx}] ---\n\n[Empty or Scanned Page]")
                except Exception as p_err:
                    logger.warning(f"pdfplumber error on page {page_idx} of '{filename}': {p_err}")
                    pages_text.append(f"--- [Page {page_idx}] ---\n\n[Page Content Unreadable]")

        if pages_text:
            return "\n\n".join(pages_text).strip()
    except Exception as e:
        logger.error(f"pdfplumber fallback also failed on '{filename}': {e}")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Corrupt or unreadable PDF file '{filename}': {str(e)}"
        )

    return "\n\n".join(pages_text).strip()


def extract_from_docx(file_bytes: bytes, filename: str) -> str:
    """Extract headings, paragraphs, and tables in document order from Word documents."""
    if not file_bytes:
        return ""

    file_stream = io.BytesIO(file_bytes)
    try:
        import docx
        doc = docx.Document(file_stream)
        elements: List[str] = []

        # Iterate through document body elements in chronological order
        for child in doc.element.body:
            tag = child.tag
            if tag.endswith("p"):
                # Paragraph element
                p = docx.text.paragraph.Paragraph(child, doc)
                text = p.text.strip()
                if text:
                    # Check if paragraph has heading style
                    style_name = p.style.name if (p.style and hasattr(p.style, "name")) else ""
                    if style_name and ("Heading" in style_name or "Title" in style_name):
                        elements.append(f"## {text}")
                    else:
                        elements.append(text)
            elif tag.endswith("tbl"):
                # Table element
                tbl = docx.table.Table(child, doc)
                table_rows = []
                for row in tbl.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        # De-duplicate contiguous identical cells from column/row spans
                        unique_cells = []
                        for c in row_cells:
                            if not unique_cells or unique_cells[-1] != c:
                                unique_cells.append(c)
                        table_rows.append(" | ".join(unique_cells))
                if table_rows:
                    elements.append("\n".join(table_rows))

        if not elements:
            # Fallback to standard paragraphs if body iteration was empty
            for p in doc.paragraphs:
                t = p.text.strip()
                if t:
                    elements.append(t)

        return "\n\n".join(elements).strip()

    except Exception as e:
        err_msg = str(e)
        logger.error(f"Error parsing Word document '{filename}': {err_msg}")
        if "Package not found" in err_msg or file_bytes.startswith(b"\xd0\xcf\x11\xe0"):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Legacy binary Word file (.doc) is not supported directly. Please convert '{filename}' to modern Word format (.docx)."
            )
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Corrupt or unreadable Word document '{filename}': {err_msg}"
        )


def extract_from_pptx(file_bytes: bytes, filename: str) -> str:
    """Extract slide titles, bullet points, text frames, tables, and speaker notes from PowerPoint decks."""
    if not file_bytes:
        return ""

    file_stream = io.BytesIO(file_bytes)
    try:
        import pptx
        prs = pptx.Presentation(file_stream)
        slides_text: List[str] = []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            # 1. Determine slide title
            slide_title = ""
            title_shape = None
            try:
                if slide.shapes.title and slide.shapes.title.text:
                    slide_title = slide.shapes.title.text.strip()
                    title_shape = slide.shapes.title
            except Exception:
                pass

            if not slide_title:
                # Search for first text frame that looks like a title
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        first_line = shape.text_frame.text.strip().split("\n")[0].strip()
                        if 0 < len(first_line) <= 100:
                            slide_title = first_line
                            title_shape = shape
                            break

            if not slide_title:
                slide_title = f"Slide {slide_idx}"

            slide_header = f"--- [Slide {slide_idx}: {slide_title}] ---"
            slide_content: List[str] = []

            # 2. Extract shape content (text frames, bullets, tables)
            for shape in slide.shapes:
                if shape == title_shape:
                    continue

                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        p_text = p.text.strip()
                        if p_text:
                            # Bullet indent formatting
                            indent = "  " * (p.level or 0)
                            slide_content.append(f"{indent}• {p_text}")

                elif shape.has_table:
                    table_rows = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            table_rows.append(" | ".join(row_cells))
                    if table_rows:
                        slide_content.append("\n".join(table_rows))

            # 3. Extract speaker notes if present
            try:
                if slide.has_notes_slide and slide.notes_slide:
                    notes_tf = slide.notes_slide.notes_text_frame
                    if notes_tf and notes_tf.text.strip():
                        notes_clean = notes_tf.text.strip()
                        slide_content.append(f"[Speaker Notes]: {notes_clean}")
            except Exception as notes_err:
                logger.debug(f"Could not read speaker notes for slide {slide_idx}: {notes_err}")

            body_text = "\n\n".join(slide_content).strip()
            if body_text:
                slides_text.append(f"{slide_header}\n\n{body_text}")
            else:
                slides_text.append(f"{slide_header}\n\n[Visual or Empty Slide]")

        return "\n\n".join(slides_text).strip()

    except Exception as e:
        err_msg = str(e)
        logger.error(f"Error parsing PowerPoint presentation '{filename}': {err_msg}")
        if file_bytes.startswith(b"\xd0\xcf\x11\xe0"):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Legacy binary PowerPoint file (.ppt) is not supported directly. Please convert '{filename}' to modern presentation format (.pptx)."
            )
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Corrupt or unreadable PowerPoint presentation '{filename}': {err_msg}"
        )


def extract_from_plaintext(file_bytes: bytes, filename: str) -> str:
    """Decode plain text files using utf-8 with latin-1 fallback."""
    if not file_bytes:
        return ""
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return file_bytes.decode("latin-1", errors="replace")
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unable to decode text file '{filename}': {str(e)}"
            )


def extract_from_image_guideline(file_bytes: bytes, filename: str) -> str:
    """Extract clean text and brand specifications from an image (PNG, JPG, JPEG, WEBP)
    using Pillow decoding, preprocessing, and Tesseract OCR with Gemini Vision fallback.
    """
    if len(file_bytes) > MAX_IMAGE_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"Image '{filename}' exceeds maximum allowed size of 10MB."
        )

    # 1. Pillow Safe Decoding & Verification
    try:
        from PIL import Image
        img_stream = io.BytesIO(file_bytes)
        image = Image.open(img_stream)
        image.verify()  # Verifies file integrity and header
        image = Image.open(io.BytesIO(file_bytes))  # Reopen after verify
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid or corrupted image '{filename}': {str(e)}"
        )

    # 2. Format Validation
    fmt = (image.format or "").upper()
    if fmt not in ("PNG", "JPEG", "JPG", "WEBP"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported image format '{fmt}' for '{filename}'. Supported formats: JPG, JPEG, PNG, WEBP."
        )

    # 3. Preprocessing: Bound dimensions to 1920x1080 to accelerate OCR & bound memory
    max_w, max_h = 1920, 1080
    w, h = image.size
    if w > max_w or h > max_h:
        ratio = min(max_w / max(w, 1), max_h / max(h, 1))
        new_size = (int(w * ratio), int(h * ratio))
        image = image.resize(new_size, Image.Resampling.LANCZOS)

    # Convert mode to RGB or grayscale if RGBA / Palette
    if image.mode not in ("RGB", "L"):
        image = image.convert("RGB")

    # 4. Tesseract OCR Extraction
    ocr_text = ""
    try:
        import pytesseract
        import shutil

        if not getattr(extract_from_image_guideline, "_checked_paths", False):
            sys_tess = shutil.which("tesseract")
            if sys_tess:
                pytesseract.pytesseract.tesseract_cmd = sys_tess
            else:
                for candidate in [
                    os.getenv("TESSERACT_CMD"),
                    "/usr/bin/tesseract",
                    "/usr/local/bin/tesseract",
                    "/opt/homebrew/bin/tesseract",
                    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                ]:
                    if candidate and os.path.exists(candidate):
                        pytesseract.pytesseract.tesseract_cmd = candidate
                        break
            setattr(extract_from_image_guideline, "_checked_paths", True)

        ocr_text = pytesseract.image_to_string(image)
        if ocr_text:
            ocr_text = ocr_text.strip()
    except (pytesseract.TesseractNotFoundError, FileNotFoundError):
        logger.info(f"Tesseract OCR not installed or not in PATH. Checking Gemini Vision fallback for '{filename}'.")
    except Exception as tess_err:
        logger.warning(f"Tesseract OCR warning on '{filename}': {tess_err}")

    # Evaluate OCR text quality: check for meaningful words
    cleaned_ocr_words = [w for w in re.split(r"\s+", ocr_text) if len(w) > 1 and any(c.isalnum() for c in w)]
    if len(cleaned_ocr_words) >= 5 and len(ocr_text) >= 25:
        logger.info(f"Tesseract OCR extracted {len(cleaned_ocr_words)} words from '{filename}'.")
        return ocr_text

    # 5. Multimodal Fallback: Gemini Vision if configured
    try:
        from app.core.config import settings
        from app.services.extractors import extract_text_via_gemini_vision
        gemini_key = os.getenv("GEMINI_API_KEY", "") or getattr(settings, "GEMINI_API_KEY", "")
        if (os.environ.get("TESTING") == "1") or (gemini_key and gemini_key.strip() and gemini_key != "your_gemini_api_key_here"):
            vision_text = extract_text_via_gemini_vision(file_bytes, (gemini_key or "mock_test_key").strip())
            if vision_text and len(vision_text.strip()) >= 20:
                lower_vision = vision_text.lower()
                if any(phrase in lower_vision for phrase in (
                    "contains no text",
                    "no visible text",
                    "no text visible",
                    "no readable text",
                    "does not contain any text",
                    "is a solid",
                    "blank image",
                )):
                    logger.info(f"Gemini Vision determined that '{filename}' contains no readable brand text.")
                    return f"[Document '{filename}' contained no readable text or consists entirely of images/scans.]"

                logger.info(f"Gemini Vision successfully transcribed brand text from '{filename}'.")
                return vision_text.strip()
    except Exception as vision_ex:
        logger.warning(f"Gemini Vision fallback failed for '{filename}': {vision_ex}")

    # 6. If OCR produced any partial text, return it
    if ocr_text and len(ocr_text.strip()) >= 15:
        return ocr_text.strip()

    # Empty or unreadable
    logger.warning(f"Image '{filename}' parsed successfully but produced empty or unreadable text.")
    return f"[Document '{filename}' contained no readable text or consists entirely of images/scans.]"


def extract_text_from_upload(file_bytes: bytes, filename: str) -> str:
    """Unified entrypoint for extracting clean text, hierarchy, and metadata
    from enterprise uploads (PDF, DOCX, PPTX, TXT/MD/CSV/JSON, and Images: JPG, JPEG, PNG, WEBP).
    """
    if not file_bytes:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file is empty (0 bytes)."
        )

    clean_filename = (filename or "upload").strip()
    ext = os.path.splitext(clean_filename)[1].lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file type. Please upload a PDF (.pdf), Word (.docx), PowerPoint (.pptx), Text (.txt), or Image file (.jpg, .jpeg, .png, .webp)."
        )

    extracted_raw = ""

    if ext == ".pdf":
        extracted_raw = extract_from_pdf(file_bytes, clean_filename)
    elif ext in (".docx", ".doc"):
        extracted_raw = extract_from_docx(file_bytes, clean_filename)
    elif ext in (".pptx", ".ppt"):
        extracted_raw = extract_from_pptx(file_bytes, clean_filename)
    elif ext in IMAGE_EXTENSIONS:
        extracted_raw = extract_from_image_guideline(file_bytes, clean_filename)
    else:
        # Plain text formats: .txt, .md, .csv, .json
        extracted_raw = extract_from_plaintext(file_bytes, clean_filename)

    # Validate length and truncate safely if exceeding enterprise LLM context limit
    validated_text = validate_and_truncate_text(extracted_raw, max_words=MAX_WORD_LIMIT)

    if not validated_text.strip():
        logger.warning(f"File '{clean_filename}' parsed successfully but produced empty text.")
        return f"[Document '{clean_filename}' contained no readable text or consists entirely of images/scans.]"

    return validated_text.strip()
