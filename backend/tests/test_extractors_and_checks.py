"""Unit and integration tests for extractors, deterministic checks, and guidelines API."""
import io
import os
import sys
from datetime import datetime, timezone, timedelta
import pytest
from PIL import Image, ImageDraw
import docx
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

# Ensure backend root is on sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.core.database import Base, get_db
from app.core.security import create_access_token, get_password_hash
from app.main import app
from app.models.brand import Brand
from app.models.guideline import GuidelineRule, RuleCategory
from app.models.user import User, UserRole
from app.services.deterministic_checks import (
    calculate_contrast_ratio,
    calculate_relative_luminance,
    check_banned_terms,
    check_color_contrast,
    check_expired_dates,
    check_expired_dates_and_promotions,
    check_html_color_contrast,
    check_image_accessibility,
    check_missing_alt_text,
    evaluate_contrast_compliance,
    run_all_deterministic_checks,
)
from app.services.extractors import (
    extract_from_docx,
    extract_from_image,
    extract_from_pdf,
    extract_from_url,
    validate_url_for_ssrf,
)
from tests.conftest import TestingSessionLocal, client



# ============================================================================
# EXTRACTORS TESTS
# ============================================================================

def test_extract_from_docx():
    # Create docx in memory
    doc = docx.Document()
    doc.add_heading("Brand Voice Guidelines", level=1)
    doc.add_paragraph("Always speak with conviction and clarity.")
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Allowed"
    table.cell(0, 1).text = "Disallowed"
    table.cell(1, 0).text = "Forward-thinking"
    table.cell(1, 1).text = "Best in the world"

    docx_bytes = io.BytesIO()
    doc.save(docx_bytes)
    docx_bytes.seek(0)

    extracted = extract_from_docx(docx_bytes.getvalue())
    assert "Brand Voice Guidelines" in extracted
    assert "Always speak with conviction and clarity." in extracted
    assert "Forward-thinking" in extracted
    assert "Best in the world" in extracted


def test_extract_from_image_fallback():
    # Create a small RGB image in memory
    img = Image.new("RGB", (100, 50), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((10, 10), "Brand", fill=(0, 0, 0))

    img_bytes = io.BytesIO()
    img.save(img_bytes, format="PNG")
    img_bytes.seek(0)

    # Should execute without throwing unhandled exceptions
    result = extract_from_image(img_bytes.getvalue())
    assert isinstance(result, str)


def test_ssrf_protection():
    # Private IP addresses must be rejected
    with pytest.raises(ValueError, match="blocked"):
        validate_url_for_ssrf("http://127.0.0.1/admin")

    with pytest.raises(ValueError, match="blocked"):
        validate_url_for_ssrf("http://localhost:3000")

    with pytest.raises(ValueError, match="blocked"):
        validate_url_for_ssrf("http://10.0.0.1/internal")

    with pytest.raises(ValueError, match="blocked"):
        validate_url_for_ssrf("http://192.168.1.1/router")

    with pytest.raises(ValueError, match="blocked"):
        validate_url_for_ssrf("http://169.254.169.254/latest/meta-data")

    # Invalid schemes must be rejected
    with pytest.raises(ValueError, match="scheme"):
        validate_url_for_ssrf("ftp://example.com/file.txt")

    with pytest.raises(ValueError, match="scheme"):
        validate_url_for_ssrf("file:///etc/passwd")


# ============================================================================
# DETERMINISTIC CHECKS TESTS
# ============================================================================

def test_check_banned_terms():
    sample_text = "Our software offers unbeatable prices and is a miracle cure for workflow bottlenecks."
    
    findings = check_banned_terms(sample_text)
    flagged_terms = [f["flagged_text"].lower() for f in findings]
    
    assert "unbeatable prices" in flagged_terms
    assert "miracle cure" in flagged_terms
    assert any(f["severity"] == "high" for f in findings)
    assert all(f.get("is_deterministic") is True for f in findings)

    # Test with explicit prohibited_terms list
    direct_findings = check_banned_terms(
        "Our product outperforms CompetitorZ in speed.",
        prohibited_terms=["CompetitorZ"]
    )
    assert len(direct_findings) == 1
    assert direct_findings[0]["flagged_text"] == "CompetitorZ"
    assert direct_findings[0]["is_deterministic"] is True


def test_check_banned_terms_with_custom_rules_and_competitors():
    custom_rule = {
        "id": "rule-proh-1",
        "category": "prohibited_claim",
        "title": "Quantum Guarantee",
        "rule_text": "Do not claim quantum computing supremacy without peer review.",
    }
    sample_text = "We have achieved Quantum Guarantee and outperformed CompetitorX in latency."

    findings = check_banned_terms(
        content=sample_text,
        prohibited_rules=[custom_rule],
        competitors=["CompetitorX"]
    )
    flagged = [f["flagged_text"] for f in findings]
    assert "Quantum Guarantee" in flagged
    assert "CompetitorX" in flagged
    assert any(f["severity"] == "critical" for f in findings)


def test_check_expired_dates():
    # Expired date in the past
    past_text = "Summer sale offer ends on October 15, 2022. Hurry before it's gone!"
    findings = check_expired_dates(past_text)
    assert len(findings) >= 1
    assert findings[0]["issue_type"] == "Expired Promotion"
    assert "October 15, 2022" in findings[0]["explanation"]
    assert findings[0]["is_deterministic"] is True

    # Future date should not trigger expired finding
    future_date_str = (datetime.now(timezone.utc) + timedelta(days=365)).strftime("%B %d, %Y")
    future_text = f"Annual promotion valid until {future_date_str}."
    future_findings = check_expired_dates(future_text)
    assert len(future_findings) == 0


def test_check_image_accessibility_and_missing_alt_text():
    # 1. Missing alt
    images = [
        {"src": "hero.jpg", "alt": None},
        {"src": "team.png", "alt": ""},
        {"src": "feature.png", "alt": "photo"},
        {"src": "diagram.svg", "alt": "Architecture diagram illustrating data flow"}
    ]
    findings = check_missing_alt_text(images)
    issue_types = [f["issue_type"] for f in findings]

    assert "Accessibility: Missing Alt Text" in issue_types
    assert "Accessibility: Empty Alt Text" in issue_types
    assert "Accessibility: Generic Alt Text" in issue_types
    assert all(f.get("is_deterministic") is True for f in findings)
    # The 4th image is fully accessible and should not have a finding
    assert len(findings) == 3


def test_color_contrast_utility():
    # Black on White: 21:1
    res_bw = check_color_contrast("#000000", "#FFFFFF")
    assert res_bw["contrast_ratio"] == 21.0
    assert res_bw["passes_aa"] is True
    assert res_bw["passes_aaa"] is True
    assert res_bw["is_deterministic"] is True

    # Low contrast: Light Gray on White (#CCCCCC on #FFFFFF)
    res_low = check_color_contrast("#CCCCCC", "#FFFFFF")
    assert res_low["contrast_ratio"] < 4.5
    assert res_low["passes_aa"] is False
    assert res_low["is_deterministic"] is True

    # HTML inline style scanner
    html_sample = '<div style="color: #CCCCCC; background-color: #FFFFFF;">Low contrast text</div>'
    contrast_findings = check_html_color_contrast(html_sample)
    assert len(contrast_findings) >= 1
    assert "Low Color Contrast" in contrast_findings[0]["issue_type"]
    assert contrast_findings[0]["is_deterministic"] is True


def test_combined_deterministic_runner():
    content = "Get our miracle cure! Valid through January 10, 2021."
    html = '<img src="promo.png"><div style="color: #EEEEEE; background: #FFFFFF;">Text</div>'
    
    findings = run_all_deterministic_checks(content=content, raw_html=html)
    assert len(findings) >= 3
    types = [f["issue_type"] for f in findings]
    assert any("Prohibited Claim" in t for t in types)
    assert any("Expired Promotion" in t for t in types)
    assert any("Missing Alt Text" in t for t in types)



# ============================================================================
# GUIDELINES INGESTION API TESTS
# ============================================================================

def test_guidelines_api_ingestion_and_retrieval():
    db = TestingSessionLocal()
    # Create test user and brand
    user = User(
        email="guideline_mgr@acme.com",
        hashed_password=get_password_hash("SecretAcmePass!"),
        role=UserRole.BrandManager
    )
    db.add(user)
    db.commit()
    db.refresh(user)

    brand = Brand(
        name="Nexus AI",
        description="Next-gen robotics",
        created_by=user.id
    )
    db.add(brand)
    db.commit()
    db.refresh(brand)
    
    user_id = str(user.id)
    user_email = user.email
    brand_id = str(brand.id)
    db.close()

    # Generate JWT for request
    token = create_access_token({"sub": user_id, "email": user_email, "role": "BrandManager"})
    headers = {"Authorization": f"Bearer {token}"}


    # 1. Ingest guidelines via raw text
    guideline_raw_text = """
1. Tone of Voice: Always communicate with optimism, technical precision, and brevity.
2. Prohibited Claims: Never claim zero risk or guaranteed 100% cure in product copy.
3. Visual Guidelines: Primary color hex code #0055FF must always meet 4.5:1 contrast against backgrounds.
4. Accessibility Standards: Every visual element and diagram requires descriptive screen reader alt text.
"""
    post_res = client.post(
        "/api/guidelines",
        data={
            "brand_id": brand_id,
            "raw_text": guideline_raw_text,
        },
        headers=headers
    )
    assert post_res.status_code == 201, post_res.text
    rules = post_res.json()
    assert len(rules) >= 4
    categories = [r["category"] for r in rules]
    assert "tone" in categories
    assert "prohibited_claim" in categories
    assert "visual" in categories
    assert "accessibility" in categories

    # 2. Fetch guidelines via GET /api/guidelines
    get_res = client.get(f"/api/guidelines?brand_id={brand_id}", headers=headers)
    assert get_res.status_code == 200
    fetched_rules = get_res.json()
    assert len(fetched_rules) >= 4

    # 3. Filter by category
    filter_res = client.get(
        f"/api/guidelines?brand_id={brand_id}&category=prohibited_claim",
        headers=headers
    )
    assert filter_res.status_code == 200
    filtered_rules = filter_res.json()
    assert len(filtered_rules) >= 1
    assert all(r["category"] == "prohibited_claim" for r in filtered_rules)

    # 4. Direct structured rule creation via POST /api/guidelines
    struct_res = client.post(
        "/api/guidelines",
        data={
            "brand_id": brand_id,
            "category": "fact",
            "title": "Clinical Trial Metric",
            "rule_text": "All effectiveness claims must cite the 2026 Phase III clinical trial.",
            "rule_code": "FACT-100"
        },
        headers=headers
    )
    assert struct_res.status_code == 201
    created_rule = struct_res.json()
    assert created_rule["title"] == "Clinical Trial Metric"
    assert created_rule["category"] == "fact"

    # 5. Upload document via POST /api/guidelines/upload
    doc_content = b"Rule 1: Always protect user data.\nRule 2: Never disclose unreleased features."
    upload_res = client.post(
        "/api/guidelines/upload",
        data={"brand_id": brand_id},
        files={"file": ("security_guidelines.txt", doc_content, "text/plain")},
        headers=headers
    )
    assert upload_res.status_code == 201
    uploaded_rules = upload_res.json()
    assert len(uploaded_rules) >= 2

