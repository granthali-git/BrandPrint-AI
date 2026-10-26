"""Automated verification suite for universal document parsing pipeline (PDF, DOCX, PPTX, TXT)."""
import io
import sys
import os

# Add backend directory to path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fastapi import HTTPException
from app.services.document_parser import (
    extract_text_from_upload,
    extract_from_pdf,
    extract_from_docx,
    extract_from_pptx,
    extract_from_plaintext
)


def create_sample_pdf() -> bytes:
    """Generate a multi-page PDF in-memory using reportlab or pypdf."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=letter)
        # Page 1
        c.drawString(100, 750, "AstraForge Cloud Executive Summary")
        c.drawString(100, 720, "Our infrastructure provides 99.95% verified uptime backed by real-time status.")
        c.showPage()
        # Page 2
        c.drawString(100, 750, "Compliance and Security Controls")
        c.drawString(100, 720, "Audited under SOC-2 Type II standards with verified compliance logs.")
        c.showPage()
        c.save()
        buffer.seek(0)
        return buffer.getvalue()
    except ImportError:
        # Fallback to minimal raw PDF with 2 pages if reportlab is not installed
        import pypdf
        writer = pypdf.PdfWriter()
        p1 = writer.add_blank_page(width=612, height=792)
        p2 = writer.add_blank_page(width=612, height=792)
        buffer = io.BytesIO()
        writer.write(buffer)
        buffer.seek(0)
        return buffer.getvalue()


def create_sample_docx() -> bytes:
    """Generate a DOCX file in-memory with headings, paragraphs, and tables."""
    import docx
    doc = docx.Document()
    doc.add_heading("Product Specification & Compliance", level=1)
    doc.add_paragraph("AstraForge Cloud offers enterprise multi-tenant container orchestration.")
    
    # Add a table
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Tier"
    hdr_cells[1].text = "SLA"
    hdr_cells[2].text = "Security"
    
    row_cells = table.add_row().cells
    row_cells[0].text = "Enterprise"
    row_cells[1].text = "99.95% High Availability"
    row_cells[2].text = "SOC-2 Type II Certified"

    doc.add_paragraph("Contact our enterprise solutions architecture team for custom onboarding.")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def create_sample_pptx() -> bytes:
    """Generate a 3-slide PPTX deck with titles, bullet points, tables, and speaker notes."""
    import pptx
    from pptx.util import Inches, Pt

    prs = pptx.Presentation()
    blank_layout = prs.slide_layouts[6]  # Blank slide layout

    # Slide 1: Executive Overview
    s1 = prs.slides.add_slide(blank_layout)
    txBox = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Executive Overview"
    p.font.size = Pt(28)

    txBox2 = s1.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(3))
    tf2 = txBox2.text_frame
    p1 = tf2.paragraphs[0]
    p1.text = "Enterprise cloud scalability verified across 10 global regions"
    p1.level = 0
    p2 = tf2.add_paragraph()
    p2.text = "Zero downtime architecture with automated failover"
    p2.level = 1

    # Speaker notes for Slide 1
    notes_slide = s1.notes_slide
    notes_slide.notes_text_frame.text = "Emphasize multi-region verified infrastructure to stakeholders."

    # Slide 2: Security & Governance
    s2 = prs.slides.add_slide(blank_layout)
    txBox3 = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    tf3 = txBox3.text_frame
    p3 = tf3.paragraphs[0]
    p3.text = "Security & Governance"
    p3.font.size = Pt(28)

    # Add a table on Slide 2
    table_shape = s2.shapes.add_table(2, 2, Inches(1), Inches(2.5), Inches(6), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Audit Standard"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "SOC-2 Type II"
    table.cell(1, 1).text = "Active & Verified"

    # Slide 3: Roadmap & SLAs
    s3 = prs.slides.add_slide(blank_layout)
    txBox4 = s3.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    tf4 = txBox4.text_frame
    p4 = tf4.paragraphs[0]
    p4.text = "Roadmap & SLAs"
    p4.font.size = Pt(28)

    txBox5 = s3.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(2))
    tf5 = txBox5.text_frame
    p5 = tf5.paragraphs[0]
    p5.text = "Guaranteed 99.95% verified uptime with transparent public dashboards"

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def run_all_tests():
    print("================================================================================")
    print("UNIVERSAL DOCUMENT EXTRACTOR VERIFICATION (PDF, DOCX, PPTX, TXT)")
    print("================================================================================")

    # Test 1: PowerPoint (.pptx) 3-Slide Deck
    print("\n[TEST 1] Testing 3-Slide PowerPoint Deck Extraction (.pptx):")
    pptx_bytes = create_sample_pptx()
    pptx_text = extract_text_from_upload(pptx_bytes, "quarterly_brand_deck.pptx")
    print(pptx_text)

    assert "--- [Slide 1: Executive Overview] ---" in pptx_text, "Missing Slide 1 delimiter"
    assert "--- [Slide 2: Security & Governance] ---" in pptx_text, "Missing Slide 2 delimiter"
    assert "--- [Slide 3: Roadmap & SLAs] ---" in pptx_text, "Missing Slide 3 delimiter"
    assert "• Enterprise cloud scalability" in pptx_text, "Missing bullet point"
    assert "[Speaker Notes]: Emphasize multi-region" in pptx_text, "Missing speaker notes"
    assert "SOC-2 Type II | Active & Verified" in pptx_text, "Missing table cells"
    print("\n[PASS] PowerPoint .pptx extracted slides, bullets, tables, and notes cleanly!")

    # Test 2: Word Document (.docx)
    print("\n[TEST 2] Testing Word Document Extraction (.docx):")
    docx_bytes = create_sample_docx()
    docx_text = extract_text_from_upload(docx_bytes, "enterprise_spec.docx")
    print(docx_text)

    assert "## Product Specification & Compliance" in docx_text or "Product Specification & Compliance" in docx_text
    assert "Tier | SLA | Security" in docx_text
    assert "Enterprise | 99.95% High Availability | SOC-2 Type II Certified" in docx_text
    print("\n[PASS] Word .docx extracted headings, paragraphs, and tables cleanly!")

    # Test 3: Multi-Page PDF (.pdf)
    print("\n[TEST 3] Testing Multi-Page PDF Extraction (.pdf):")
    pdf_bytes = create_sample_pdf()
    pdf_text = extract_text_from_upload(pdf_bytes, "annual_report.pdf")
    print(pdf_text)

    assert "--- [Page 1] ---" in pdf_text
    assert "--- [Page 2] ---" in pdf_text
    print("\n[PASS] PDF .pdf extracted multi-page boundaries with delimiters cleanly!")

    # Test 4: Plain Text (.txt)
    print("\n[TEST 4] Testing Plaintext Extraction (.txt):")
    txt_content = "AstraForge Cloud delivers enterprise grade scalability.\nNo superlatives or hype."
    txt_bytes = txt_content.encode("utf-8")
    txt_text = extract_text_from_upload(txt_bytes, "guidelines.txt")
    assert "AstraForge Cloud delivers" in txt_text
    print("\n[PASS] Plaintext extracted successfully!")

    # Test 5: Unsupported File Extension
    print("\n[TEST 5] Testing Unsupported File Rejection:")
    try:
        extract_text_from_upload(b"MZ\x90\x00", "malware.exe")
        assert False, "Should have raised HTTPException for .exe"
    except HTTPException as e:
        assert e.status_code == 400
        assert "Unsupported file type. Please upload a PDF (.pdf), Word (.docx), PowerPoint (.pptx), or Text (.txt) file." in str(e.detail)
        print(f"[PASS] Correctly rejected with HTTP 400: {e.detail}")

    print("\n================================================================================")
    print("ALL DOCUMENT PARSER TESTS PASSED 100%!")
    print("================================================================================")


if __name__ == "__main__":
    run_all_tests()
